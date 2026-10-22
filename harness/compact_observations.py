"""Deterministic enrichment of minimal model observations; never repairs findings."""
from __future__ import annotations

import json
from pathlib import Path
import re


def _text(value, field):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f'{field} must be nonempty text')
    return value


def _object(value, allowed, required, field):
    if not isinstance(value, dict) or not required <= value.keys() or value.keys() - allowed:
        raise ValueError(f'{field}: missing or unknown fields')


def _list(value, field, nonempty=False):
    if not isinstance(value, list) or (nonempty and not value):
        raise ValueError(f'{field} must be a list' + (' with evidence' if nonempty else ''))
    return value


def _resolve(ref, root, index):
    _text(ref, 'evidence reference')
    sources = index['sources']

    def located(alias, first, last, label):
        if alias not in sources:
            raise ValueError(f'evidence source unavailable: {alias}')
        source = sources[alias]
        if first < 1 or last < first or last > source['line_count']:
            raise ValueError(f'evidence line range out of bounds: {ref}')
        return {'path': str(root/source['relative_path']),
                'locator': f'{label}; lines {first}-{last}', 'ref': ref}

    if ref == 'trace:all':
        return located('trace', 1, sources['trace']['line_count'], ref)
    if ref == 'trace:score':
        sections = index['score_sections']
        if len(sections) != 1:
            raise ValueError('trace:score requires exactly one indexed score section')
        return located('trace', sections[0]['line_start'], sections[0]['line_end'], ref)
    match = re.fullmatch(r'event:(\d+)|events:(\d+)-(\d+)', ref)
    if match:
        first = int(match[1] if match[1] is not None else match[2])
        last = int(match[1] if match[1] is not None else match[3])
        events = {item['event']: item for item in index['events']}
        if last < first or last-first+1 > len(events) or any(i not in events for i in range(first, last+1)):
            raise ValueError(f'event range missing or out of bounds: {ref}')
        selected = [events[i] for i in range(first, last+1)]
        if any(a['line_start'] >= b['line_start'] for a, b in zip(selected, selected[1:])):
            raise ValueError(f'event range is not in source order: {ref}')
        return located('trace', selected[0]['line_start'], selected[-1]['line_end'], ref)
    match = re.fullmatch(r'(task|system):lines:(\d+)(?:-(\d+))?', ref)
    if match:
        return located(match[1], int(match[2]), int(match[3] or match[2]), ref)
    match = re.fullmatch(r'input:(.+)!([A-Z]+[1-9]\d*(?::[A-Z]+[1-9]\d*)?)', ref)
    if match:
        if 'input' not in sources:
            raise ValueError('evidence source unavailable: input')
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries
        sheet, coordinates = match.groups()
        path = root/sources['input']['relative_path']
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet not in workbook.sheetnames:
                raise ValueError(f'unknown evidence worksheet: {sheet}')
            worksheet = workbook[sheet]
            min_col, min_row, max_col, max_row = range_boundaries(coordinates)
            if (min_col > max_col or min_row > max_row or
                    max_col > worksheet.max_column or max_row > worksheet.max_row):
                raise ValueError(f'workbook evidence range out of bounds: {ref}')
        finally:
            workbook.close()
        return {'path': str(path), 'locator': f'{sheet}!{coordinates}', 'ref': ref}
    raise ValueError(f'unknown evidence reference: {ref}')


def _check_score_trace_promises(result, root):
    """Enforce the mechanical score-trace promises on an attributed attempt.

    Nothing here judges a finding: it only requires the prefixes the trace makes
    decidable, so a missing reason or an uncorrected visible failure is loud.
    """
    trace_path = root / 'score-trace.json'
    record = json.loads(trace_path.read_bytes()) if trace_path.is_file() else None
    statements = [item['statement'] for item in result['observations']]
    traces = [text for text in statements if text.startswith('SCORE TRACE:')]
    rewards = [text for text in statements if text.startswith('FEEDBACK:')]
    if record is None or record.get('reconciliation', {}).get('status') != 'reproduced':
        detail = ('no score trace is supplied' if record is None
                  else 'the replay did not reproduce the receipt')
        if traces or rewards:
            raise ValueError(f'{detail}; no SCORE TRACE: or FEEDBACK: observation is allowed')
        return
    if (record.get('official', {}).get('reward') or 0) >= 1:
        if traces:
            raise ValueError('a solved attempt carries a SCORE TRACE: observation')
        return
    if len(traces) != 1:
        raise ValueError('an attributed unsolved attempt requires exactly one SCORE TRACE: observation')
    numbers = [int(value) for value in re.findall(r'REASON (\d+):', traces[0])]
    if not numbers or numbers != list(range(1, len(numbers) + 1)):
        raise ValueError('SCORE TRACE: reasons must be REASON 1:, REASON 2:, ... in order with no gaps')
    from score_trace import feedback_condition
    if feedback_condition(record) and not rewards:
        raise ValueError('the score trace shows an uncorrected visible failure; write a FEEDBACK: observation')
    if len(rewards) > 1:
        raise ValueError('at most one FEEDBACK: observation is allowed')


def enrich_observations(raw, manifest_path):
    """Return canonical script-owned metadata, preserving raw semantic values.

    Callers must retain raw JSON separately. Omission, explicit null and declared
    uncertainty remain distinguishable; an empty legacy list is not certainty.
    """
    try:
        from .worker_inputs import _read_manifest
        from .parallel_records import validate_observations
    except ImportError:
        from worker_inputs import _read_manifest
        from parallel_records import validate_observations
    manifest = _read_manifest(manifest_path)
    _object(raw, {'observations', 'limitations'}, {'observations'}, 'compact observations')
    root = Path(manifest_path).resolve().parent
    if not any(item['relative_path'] == 'evidence_index.json' for item in manifest['files']):
        raise ValueError('compact observations require an indexed input manifest')
    index = json.loads((root/'evidence_index.json').read_bytes())
    sample = manifest['sample_id']
    observations = []
    for i, item in enumerate(_list(raw['observations'], 'observations'), 1):
        _object(item, {'finding', 'evidence', 'uncertainty'}, {'finding', 'evidence'}, 'observation')
        finding = _text(item['finding'], 'finding')
        evidence = [_resolve(ref, root, index) for ref in _list(item['evidence'], 'evidence', True)]
        status = 'not_declared' if 'uncertainty' not in item else 'explicit_null' if item['uncertainty'] is None else 'declared'
        uncertainty = [] if status != 'declared' else [_text(item['uncertainty'], 'uncertainty')]
        observations.append({'id': f's{sample:02d}-o{i}', 'statement': finding,
                             'evidence': evidence, 'uncertainties': uncertainty,
                             'uncertainty_status': status})
    result = {'schema_version': 'trajectory-observations-v1', 'sample_id': sample,
              'observations': observations,
              'limitations_status': 'declared' if 'limitations' in raw else 'not_declared'}
    if 'limitations' in raw:
        result['evidence_limitations'] = [_text(item, 'limitation') for item in _list(raw['limitations'], 'limitations')]
    _check_score_trace_promises(result, root)
    validate_observations(result)
    return result
