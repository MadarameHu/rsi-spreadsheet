#!/usr/bin/env python3
"""Deterministic verifier probes and admission; replaces the LLM review by default.

Every historical review rejection was a verifier that accepted a wrong workbook of a
known shape. These probes build those shapes from the candidate's own finished
reference workbook and run them through the unchanged verifier via Harbor:

  control           the reference workbook re-saved unchanged       must be accepted
  constants_only    target cells hold the ground-truth values only  must be rejected
  literal_formula   target cells hold ="value" style constant formulas  must be rejected
  unscored_outputs  cells the reference wrote outside answer_position
                    are reverted to the input                       must be rejected

constants_only / literal_formula apply only when the reference itself writes formulas
into the target cells; unscored_outputs only when such cells exist. A rejected control
means the probe harness cannot judge this task: blocked, never accepted.
No model calls. Admission also needs the reference run to pass, done by the caller.
"""
from __future__ import annotations

import argparse
from concurrent.futures import ThreadPoolExecutor
from copy import copy
from datetime import date, datetime, time as dtime
import json
import math
from pathlib import Path
import sys
import tomllib

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

sys.path.insert(0, str(Path(__file__).resolve().parent))

SCHEMA = 'template-gate-v1'
NEGATIVE = ('constants_only', 'literal_formula', 'unscored_outputs')


def _formula(value):
    if hasattr(value, 'text'):  # ArrayFormula
        value = value.text
    return value if isinstance(value, str) and value.startswith('=') and len(value) > 1 else None


def _ranges(spec):
    out = []
    for item in spec['target_ranges']:
        c1, r1, c2, r2 = range_boundaries(item['range'])
        out.append((item['sheet'], c1, r1, c2, r2))
    return out


def _target_cells(spec):
    return [(s, f'{get_column_letter(c)}{r}') for s, c1, r1, c2, r2 in _ranges(spec)
            for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


def _in_targets(ranges, sheet, row, col):
    return any(s == sheet and c1 <= col <= c2 and r1 <= row <= r2 for s, c1, r1, c2, r2 in ranges)


def _norm(value):
    if isinstance(value, bool) or value is None:
        return value
    if isinstance(value, (int, float)):
        return round(float(value), 9)
    if isinstance(value, str):
        return value.strip() or None
    return value


def _literal(value):
    """A formula that evaluates to the constant, so a formula-presence check passes."""
    if value is None:
        return '=""'
    if isinstance(value, bool):
        return '=TRUE()' if value else '=FALSE()'
    if isinstance(value, (int, float)) and math.isfinite(value):
        return '=' + repr(value)
    if isinstance(value, datetime):
        t = value.time()
        return f'=DATE({value.year},{value.month},{value.day})+TIME({t.hour},{t.minute},{t.second})'
    if isinstance(value, date):
        return f'=DATE({value.year},{value.month},{value.day})'
    if isinstance(value, dtime):
        return f'=TIME({value.hour},{value.minute},{value.second})'
    if isinstance(value, str) and value == '#N/A':
        return '=NA()'
    return '="' + str(value).replace('"', '""') + '"'


def unscored_cells(input_xlsx, reference_xlsx, spec):
    """Cells the reference changed relative to the input, outside every target range."""
    ranges = _ranges(spec)
    src = openpyxl.load_workbook(input_xlsx)
    ref = openpyxl.load_workbook(reference_xlsx)
    changed = []
    try:
        for ws in ref.worksheets:
            old = src[ws.title] if ws.title in src.sheetnames else None
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None or _in_targets(ranges, ws.title, cell.row, cell.column):
                        continue
                    before = old[cell.coordinate].value if old is not None else None
                    a, b = _formula(before), _formula(cell.value)
                    same = (a == b) if (a or b) else _norm(before) == _norm(cell.value)
                    if not same:
                        changed.append((ws.title, cell.coordinate))
    finally:
        src.close()
        ref.close()
    return changed


def build(candidate, reference_xlsx, out_dir):
    """Write the probe workbooks; return plans in the shape stage_audit.run_probe expects."""
    candidate, out_dir = Path(candidate), Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=False)
    tests = candidate / 'tests'
    spec = json.loads((tests / 'task_spec.json').read_text())
    artifacts = tomllib.loads((candidate / 'task.toml').read_text()).get('artifacts') or []
    if len(artifacts) != 1:
        raise ValueError('template probes need exactly one declared workbook artifact')
    target_path = artifacts[0]
    cells = _target_cells(spec)
    truth_wb = openpyxl.load_workbook(tests / 'target.xlsx', data_only=True)
    truth = {(s, a): truth_wb[s][a].value for s, a in cells}
    truth_wb.close()

    def variant(name, edit):
        wb = openpyxl.load_workbook(reference_xlsx)
        try:
            edit(wb)
            path = out_dir / f'{name}.xlsx'
            wb.save(path)
        finally:
            wb.close()
        return path

    ref = openpyxl.load_workbook(reference_xlsx)
    formula_cells = [(s, a) for s, a in cells if s in ref.sheetnames and _formula(ref[s][a].value)]
    ref.close()
    changed = unscored_cells(candidate / 'environment/input.xlsx', reference_xlsx, spec)
    base = {'evidence': [{'path': str(reference_xlsx), 'locator': 'finished reference workbook'}],
            'target_path': target_path}
    plans = [{**base, 'id': 'control', 'expected_valid': True,
              'reason': 'Unchanged reference workbook; the verifier must accept it.',
              'workbook': str(variant('control', lambda wb: None))}]
    skipped = {}
    if formula_cells:
        def constants(wb):
            for s, a in formula_cells:
                wb[s][a].value = truth[(s, a)]

        def literals(wb):
            for s, a in formula_cells:
                wb[s][a].value = _literal(truth[(s, a)])
        plans.append({**base, 'id': 'constants_only', 'expected_valid': False,
                      'reason': f'{len(formula_cells)} target formulas replaced by their correct values.',
                      'workbook': str(variant('constants_only', constants))})
        plans.append({**base, 'id': 'literal_formula', 'expected_valid': False,
                      'reason': f'{len(formula_cells)} target formulas replaced by constant formulas.',
                      'workbook': str(variant('literal_formula', literals))})
    else:
        skipped['constants_only'] = skipped['literal_formula'] = 'reference writes no target formulas'
    if changed:
        src = openpyxl.load_workbook(candidate / 'environment/input.xlsx')
        original = {(s, a): (src[s][a].value if s in src.sheetnames else None) for s, a in changed}
        src.close()

        def revert(wb):
            for (s, a), value in original.items():
                wb[s][a].value = copy(value)
        plans.append({**base, 'id': 'unscored_outputs', 'expected_valid': False,
                      'reason': f'{len(changed)} reference outputs outside answer_position reverted, e.g. '
                                + ', '.join(f'{s}!{a}' for s, a in changed[:5]),
                      'workbook': str(variant('unscored_outputs', revert))})
    else:
        skipped['unscored_outputs'] = 'reference writes nothing outside answer_position'
    return plans, skipped


def judge(receipts, skipped):
    """Admission from executed probes: accepted / rejected / blocked, with issues."""
    issues, status = [], 'accepted'
    for r in receipts:
        if r['execution_status'] != 'completed':
            issues.append({'id': r['id'], 'verdict': 'fault', 'reason': 'probe did not finish'})
            status = 'blocked'
        elif r['actual_accept'] is not r['expected_valid']:
            if r['id'] == 'control':
                issues.append({'id': 'control', 'verdict': 'fault',
                               'reason': 'verifier rejected the unchanged reference; probes cannot judge'})
                status = 'blocked'
            else:
                issues.append({'id': r['id'], 'verdict': 'incorrect',
                               'reason': 'verifier accepted a wrong workbook: ' + r['reason']})
                if status == 'accepted':
                    status = 'rejected'
    return {'status': status, 'eligible_for_training': status == 'accepted', 'issues': issues,
            'not_applicable': skipped}


def execute(hb, cfg, benchmark, candidate, reference_xlsx, out_dir, *, max_concurrent=4):
    """Build, freeze and run every probe through the unchanged verifier; never retried."""
    import stage_audit
    out_dir = Path(out_dir)
    plans, skipped = build(candidate, reference_xlsx, out_dir / 'workbooks')
    plan_path = out_dir / 'plan.json'
    plan_path.write_text(json.dumps({'schema_version': SCHEMA, 'probes': plans}, indent=1) + '\n')
    frozen = stage_audit.freeze_probe_inputs(plans, plan_path, out_dir / 'frozen')
    with ThreadPoolExecutor(max_workers=max(1, min(max_concurrent, len(plans)))) as pool:
        receipts = list(pool.map(lambda p: stage_audit.run_probe(
            hb, cfg, benchmark, candidate, p, out_dir / 'runs' / p['id'],
            frozen['probes'][p['id']], frozen['plan_sha256']), plans))
    return {'schema_version': SCHEMA, 'origin': 'template_probes', 'llm_review': 'disabled',
            **judge(receipts, skipped), 'probes': receipts}


def reference_workbook(ep, validation=None):
    """The finished reference run's collected workbook artifact."""
    ep = Path(ep)
    validation = validation or json.loads((ep / 'diagnostics/validation.json').read_text())
    result = (validation.get('oracle') or {}).get('result_path')
    if not result:
        raise ValueError('reference run has no result path')
    candidate = Path(json.loads((ep / 'diagnostics/candidate_manifest.json').read_text())['snapshot_path']) \
        if (ep / 'diagnostics/candidate_manifest.json').exists() else ep / 'diagnostics/candidate'
    target = tomllib.loads((candidate / 'task.toml').read_text())['artifacts'][0]
    path = Path(result).parent / 'artifacts' / target.lstrip('/')
    if not path.is_file():
        raise ValueError(f'reference workbook artifact missing: {path}')
    return path


def run_gate(hb, cfg, benchmark, ep, *, max_concurrent=4):
    """Write diagnostics/template-gate.json for the committed candidate; reuse a finished one."""
    ep = Path(ep)
    receipt_path = ep / 'diagnostics/template-gate.json'
    if receipt_path.exists():
        return json.loads(receipt_path.read_text())
    candidate = ep / 'diagnostics/candidate'
    out = ep / 'diagnostics/template-probes'
    if out.exists():
        raise ValueError(f'unfinished template probes at {out}; inspect before rerunning')
    reference = reference_workbook(ep)
    result = execute(hb, cfg, benchmark, candidate, reference, out, max_concurrent=max_concurrent)
    result.update(episode=str(ep), reference_workbook=str(reference))
    receipt_path.write_text(json.dumps(result, indent=1, ensure_ascii=False) + '\n')
    return result


REVIEW_MODES = ('template', 'llm')


def trainee_runtime(cfg):
    """Trainee attempts run locally through Harbor in this distribution."""
    return 'local'


def review_mode(cfg):
    """training_pipeline.review: 'template' (default, no model) or 'llm' (opus audit)."""
    mode = (cfg.get('training_pipeline') or {}).get('review', 'template')
    if mode not in REVIEW_MODES:
        raise ValueError(f'training_pipeline.review must be one of {REVIEW_MODES}')
    return mode


def template_admission(ep):
    """Controller-side recheck before trainee calls; missing receipts are blocked."""
    leak = Path(ep) / 'diagnostics/leak-check.json'
    if not leak.is_file() or json.loads(leak.read_text()).get('status') != 'accepted':
        return {'status': 'blocked', 'issues': [{'id': 'leak_check', 'verdict': 'missing',
                                                 'reason': 'no accepted leak check receipt'}]}
    path = Path(ep) / 'diagnostics/template-gate.json'
    if not path.is_file():
        return {'status': 'blocked', 'issues': [{'id': 'template_gate', 'verdict': 'missing',
                                                 'reason': 'no template gate receipt'}]}
    return json.loads(path.read_text())


def write_route(ep, validation, cfg):
    result = route(validation, cfg)
    (Path(ep) / 'diagnostics/route.json').write_text(json.dumps(result, indent=1) + '\n')
    return result


def route(validation, cfg):
    """Training route from measured trainee attempts: rl inside the band, otherwise sft."""
    n, k = validation.get('completed_attempts') or 0, validation.get('solved_attempts') or 0
    if n == 0:
        return {'route': None, 'reason': 'no completed trainee attempts'}
    lo, hi = float(cfg['score']['band_lo']), float(cfg['score']['band_hi'])
    rate = k / n
    return {'solved': k, 'completed': n, 'pass_rate': rate, 'band': [lo, hi],
            'route': 'rl' if lo <= rate <= hi else 'sft'}


def main(argv=None):
    ap = argparse.ArgumentParser(description='run template probes on finished episodes')
    ap.add_argument('--config', required=True)
    ap.add_argument('--benchmark', default='spreadsheet')
    ap.add_argument('--out-root', required=True, help='write probe runs here, not into the episode')
    ap.add_argument('--concurrency', type=int, default=4)
    ap.add_argument('episodes', nargs='+')
    args = ap.parse_args(argv)
    from common import load_config, make_harbor
    cfg = load_config(args.config)
    hb = make_harbor(cfg, cfg['score']['setup_timeout_mult'])

    def one(ep):
        ep = Path(ep).resolve()
        name = '__'.join(ep.parts[-4:])
        try:
            r = execute(hb, cfg, args.benchmark, ep / 'diagnostics/candidate', reference_workbook(ep),
                        Path(args.out_root) / name, max_concurrent=args.concurrency)
        except Exception as exc:
            r = {'status': 'error', 'error': f'{type(exc).__name__}: {exc}'}
        (Path(args.out_root) / f'{name}.json').write_text(json.dumps(r, indent=1, ensure_ascii=False) + '\n')
        return name, r.get('status'), [i['id'] for i in r.get('issues', [])]
    with ThreadPoolExecutor(max_workers=len(args.episodes)) as pool:
        for row in pool.map(one, args.episodes):
            print(json.dumps(row), flush=True)


if __name__ == '__main__':
    main()
