"""Score trace replays recorded edits mechanically; no model calls, no LibreOffice."""
from pathlib import Path
import tempfile
import unittest

import openpyxl

from harness import score_trace as trace

ATTEMPT = '''# task-1 / attempt-0

Status: completed; termination: model_final; reward: 0.0

## Official score and execution summary
```json
{"official_score": {"reward": 0.0, "msg": "Value mismatch at C2: your workbook has 0"},
 "usage": {"completion_tokens": 10, "total_tokens": 30}, "duration_seconds": 1.5}
```

## Ordered events

### step 1 / event 0 / tool_call

```json
{"tool_name": "fill_formula", "execution_status": "error", "response_event_index": 1,
 "arguments": {"sheet_name": "S", "start_cell": "C2", "end_row": 3, "formula_template": "=A2+"}}
```

#### Event text
```
<tool_call>
```

### step 1 / event 1 / tool_response

```json
{"tool_name": "fill_formula", "execution_status": "error", "arguments": null}
```

#### Event text
```
Error: formula syntax invalid
```

### step 2 / event 2 / tool_call

```json
{"tool_name": "fill_formula", "execution_status": "success", "response_event_index": 3,
 "arguments": {"sheet_name": "S", "start_cell": "C2", "end_row": 3, "formula_template": "=A2*B$1*''"}}
```

#### Event text
```
<tool_call>
```

### step 2 / event 3 / tool_response

```json
{"tool_name": "fill_formula", "execution_status": "success", "arguments": null}
```

#### Event text
```
{"status":"success","samples":[{"address":"S!C2","formula":"=A2*B$1","value":0},{"address":"S!C3","formula":"=A3*B$1","value":0}]}
```
'''


class ScoreTraceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.attempt = self.root / 'attempt-0.md'
        self.attempt.write_text(ATTEMPT)
        wb = openpyxl.Workbook()
        wb.active.title = 'S'
        self.source = self.root / 'input.xlsx'
        wb.save(self.source)

    def test_parse_keeps_receipt_and_event_order(self):
        parsed = trace.parse_attempt(self.attempt)
        self.assertEqual((parsed['reward'], parsed['termination']), (0.0, 'model_final'))
        self.assertEqual([e['event'] for e in parsed['events']], [0, 1, 2, 3])
        self.assertEqual(trace._receipt_cell(parsed['msg']), ('C2', '0'))

    def test_replay_uses_only_successful_edits_and_recorded_formula(self):
        parsed = trace.parse_attempt(self.attempt)
        edits, failures = trace._replay(parsed['events'], self.source, self.root / 'out.xlsx')
        self.assertEqual((len(edits), failures), (1, []))
        ws = openpyxl.load_workbook(self.root / 'out.xlsx')['S']
        self.assertEqual((ws['C2'].value, ws['C3'].value), ('=A2*B$1', '=A3*B$1'))

    def test_feedback_flags_and_process_counts_are_mechanical(self):
        parsed = trace.parse_attempt(self.attempt)
        feedback = trace.last_feedback(parsed['events'])
        self.assertEqual((feedback['flags'], feedback['tool_calls_after']), (['all_zero', 'all_identical'], 0))
        stats = trace.process_stats(parsed)
        self.assertEqual((stats['tool_calls'], stats['failed_calls'], stats['identical_repeat_calls']), (2, 1, 0))

    def test_write_range_shapes_follow_the_native_tool(self):
        self.assertEqual(trace._grid('x', 3, 1, False), [['x'], ['x'], ['x']])
        self.assertEqual(trace._grid('x', 1, 1, True), [['x']])
        self.assertEqual(trace._grid([1, 2], 2, 1, False), [[1], [2]])
        self.assertEqual(trace._grid([1, 2], 1, 1, True), [[1, 2]])
        with self.assertRaises(ValueError):
            trace._grid([[1, 2], [3, 4]], 3, 3, False)

    def test_receipt_value_comparison(self):
        self.assertTrue(trace._same_as_receipt(0, '0'))
        self.assertTrue(trace._same_as_receipt('#VALUE!', "'#VALUE!'"))
        self.assertTrue(trace._same_as_receipt(None, 'None'))
        self.assertTrue(trace._same_as_receipt(False, 'False'))
        self.assertFalse(trace._same_as_receipt(6, '7'))


if __name__ == '__main__':
    unittest.main()
