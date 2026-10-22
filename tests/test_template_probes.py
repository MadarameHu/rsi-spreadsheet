import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'harness'))
import run_episode  # noqa: E402
import template_probes as tp  # noqa: E402


def book(path, cells):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'S'
    for ref, value in cells.items():
        ws[ref] = value
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


class TemplateProbeTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        self.cand = self.root / 'candidate'
        book(self.cand / 'environment/input.xlsx', {'A1': 2, 'A2': 3})
        book(self.cand / 'tests/target.xlsx', {'A1': 2, 'A2': 3, 'B1': 5})
        (self.cand / 'tests/task_spec.json').write_text(json.dumps(
            {'answer_position': 'S!B1', 'target_ranges': [{'sheet': 'S', 'range': 'B1'}]}))
        (self.cand / 'task.toml').write_text('artifacts = ["/app/data.xlsx"]\n')
        self.ref = self.root / 'ref.xlsx'

    def values(self, plans, name):
        wb = openpyxl.load_workbook(next(p['workbook'] for p in plans if p['id'] == name))
        return {c.coordinate: c.value for row in wb['S'].iter_rows() for c in row if c.value is not None}

    def test_formula_reference_gets_constant_probes_and_control(self):
        book(self.ref, {'A1': 2, 'A2': 3, 'B1': '=A1+A2'})
        plans, skipped = tp.build(self.cand, self.ref, self.root / 'w')
        self.assertEqual([p['id'] for p in plans], ['control', 'constants_only', 'literal_formula'])
        self.assertEqual(self.values(plans, 'control')['B1'], '=A1+A2')
        self.assertEqual(self.values(plans, 'constants_only')['B1'], 5)
        self.assertEqual(self.values(plans, 'literal_formula')['B1'], '=5')
        self.assertEqual(set(skipped), {'unscored_outputs'})
        self.assertTrue(all(p['target_path'] == '/app/data.xlsx' for p in plans))

    def test_outputs_outside_answer_position_are_reverted(self):
        book(self.ref, {'A1': 2, 'A2': 3, 'B1': 5, 'C1': 'Total'})
        plans, skipped = tp.build(self.cand, self.ref, self.root / 'w')
        self.assertEqual([p['id'] for p in plans], ['control', 'unscored_outputs'])
        self.assertNotIn('C1', self.values(plans, 'unscored_outputs'))
        self.assertEqual(set(skipped), {'constants_only', 'literal_formula'})

    def test_literals_escape_and_cover_types(self):
        self.assertEqual(tp._literal('a"b'), '="a""b"')
        self.assertEqual(tp._literal(None), '=""')
        self.assertEqual(tp._literal(True), '=TRUE()')
        self.assertEqual(tp._literal('#N/A'), '=NA()')

    def test_judge_statuses(self):
        def r(i, exp, act, status='completed'):
            return {'id': i, 'expected_valid': exp, 'actual_accept': act, 'execution_status': status, 'reason': 'x'}
        self.assertEqual(tp.judge([r('control', True, True), r('constants_only', False, False)], {})['status'], 'accepted')
        rejected = tp.judge([r('control', True, True), r('literal_formula', False, True)], {})
        self.assertEqual((rejected['status'], rejected['issues'][0]['id']), ('rejected', 'literal_formula'))
        self.assertEqual(tp.judge([r('control', True, False), r('constants_only', False, False)], {})['status'], 'blocked')
        self.assertEqual(tp.judge([r('control', True, None, 'infra_error')], {})['status'], 'blocked')

    def test_review_mode_and_route(self):
        self.assertEqual(tp.review_mode({}), 'template')
        self.assertEqual(tp.review_mode({'training_pipeline': {'review': 'llm'}}), 'llm')
        with self.assertRaises(ValueError):
            tp.review_mode({'training_pipeline': {'review': 'off'}})
        cfg = {'score': {'band_lo': 0.125, 'band_hi': 0.75}}
        self.assertEqual(tp.route({'completed_attempts': 16, 'solved_attempts': 2}, cfg)['route'], 'rl')
        self.assertEqual(tp.route({'completed_attempts': 16, 'solved_attempts': 13}, cfg)['route'], 'sft')
        self.assertEqual(tp.route({'completed_attempts': 16, 'solved_attempts': 1}, cfg)['route'], 'sft')
        self.assertIsNone(tp.route({'completed_attempts': 0}, cfg)['route'])

    def test_template_mode_never_starts_llm_audit(self):
        ep = self.root / 'ep'
        (ep / 'diagnostics').mkdir(parents=True)
        cfg = {'score': {'band_lo': 0.125, 'band_hi': 0.75}}
        phases = []

        def launch(hb, cfg, benchmark, ep, phase='all', **kw):
            phases.append(phase)
            return ({'status': 'reference_completed'} if phase == 'reference'
                    else {'status': 'completed', 'completed_attempts': 16, 'solved_attempts': 4})
        import host_validation, stage_audit, leak_check, defect_hits
        ok = patch.object(leak_check, 'run_for_episode', return_value={'status': 'accepted'})
        ok.start()
        self.addCleanup(ok.stop)
        with patch.object(host_validation, 'launch_validation', side_effect=launch), \
             patch.object(defect_hits, 'write_for_episode', return_value={}), \
             patch.object(stage_audit, 'audit_episode', side_effect=AssertionError('no LLM audit')), \
             patch.object(tp, 'run_gate', return_value={'status': 'accepted', 'eligible_for_training': True}):
            result = run_episode.validate_authored_candidate(object(), cfg, 'spreadsheet', ep)
        self.assertEqual((phases, result['status']), (['reference', 'trainee'], 'completed'))
        self.assertEqual(json.loads((ep / 'diagnostics/route.json').read_text())['route'], 'rl')
        phases.clear()
        with patch.object(host_validation, 'launch_validation', side_effect=launch), \
             patch.object(tp, 'run_gate', return_value={'status': 'rejected', 'eligible_for_training': False}):
            result = run_episode.validate_authored_candidate(object(), cfg, 'spreadsheet', ep)
        self.assertEqual((phases, result['status']), (['reference'], 'template_rejected'))
        ok.stop()
        phases.clear()
        with patch.object(leak_check, 'run_for_episode', return_value={'status': 'rejected'}), \
             patch.object(host_validation, 'launch_validation', side_effect=launch):
            result = run_episode.validate_authored_candidate(object(), cfg, 'spreadsheet', ep)
        self.assertEqual((phases, result['status']), ([], 'leak_rejected'))
        ok.start()

    def test_admission_requires_accepted_leak_check(self):
        ep = self.root / 'ep2'
        (ep / 'diagnostics').mkdir(parents=True)
        (ep / 'diagnostics/template-gate.json').write_text('{"status": "accepted"}')
        self.assertEqual(tp.template_admission(ep)['status'], 'blocked')
        (ep / 'diagnostics/leak-check.json').write_text('{"status": "accepted"}')
        self.assertEqual(tp.template_admission(ep)['status'], 'accepted')


if __name__ == '__main__':
    unittest.main()
