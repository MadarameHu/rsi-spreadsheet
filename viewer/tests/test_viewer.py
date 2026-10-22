import importlib
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / 'indexer'))
sys.path.insert(0, str(HERE.parent / 'server'))

import build_index  # noqa: E402


def write(path: Path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False))


def make_full_episode(root: Path, run: str = 'batch-1-00001-d1',
                      task: str = 'task-00001') -> Path:
    ep = root / 'runs' / run / 'spreadsheet' / task / 'ep01'
    diag = ep / 'diagnostics'
    write(diag / 'observations/index.json', {
        'status': 'completed',
        'observation_ids': ['s00-o1', 's00-o2'],
        'workers': [{'sample_id': 0}, {'sample_id': 1}],
        'reuse_observations_from': '/x',
    })
    write(diag / 'attribution.json', {
        'derivation': 'auto',
        'claims': [{'id': 'c1', 'evidence_strength': 'replay_verified'}],
        'training_goals': [{'id': 'g1', 'ability': 'keep references stable'}],
    })
    write(diag / 'design.json', {
        'primary_goal_id': 'c1',
        'selection_reason': 'because',
        'targets': [{'claim_id': 'c1'}],
        'preserves': ['layout'],
        'simplifies': [],
    })
    write(diag / 'leak-check.json', {
        'status': 'accepted', 'model': 'm',
        'verdicts': [{'verdict': 'pass'}], 'attempts': 1,
    })
    write(diag / 'validation.json', {
        'status': 'reference_completed',
        'oracle': {'status': 'completed', 'reward': 1.0},
        'requested_attempts': 2, 'solved_attempts': 1, 'completed_attempts': 2,
        'attempts': [
            {'id': 'a0', 'kind': 'trainee', 'status': 'completed', 'reward': 1.0},
            {'id': 'a1', 'kind': 'trainee', 'status': 'completed', 'reward': 0.0},
        ],
    })
    write(diag / 'template-gate.json', {
        'status': 'accepted', 'eligible_for_training': True,
        'probes': [{'id': 'control', 'expected_valid': True,
                    'actual_accept': True, 'execution_status': 'completed'}],
        'issues': [], 'llm_review': 'disabled',
    })
    write(diag / 'defect-hits.json', {
        'off_target': False, 'attempts': 2, 'failed': 1,
        'class_hits': {'c1': 1}, 'categories': {'hit': 1}, 'threshold': 0.8,
        'rows': [{'id': 'a1', 'reward': 0.0, 'category': 'hit', 'classes': ['c1']}],
    })
    write(diag / 'route.json', {
        'route': 'rl', 'pass_rate': 0.5, 'band': [0.125, 0.75],
        'solved': 1, 'completed': 2,
    })
    (diag / 'events.jsonl').write_text(
        json.dumps({'event': 'stage_started', 'stage': 'design'}) + '\n'
        + json.dumps({'event': 'stage_finished', 'stage': 'design'}) + '\n')
    write(ep / 'synthesis.json', {'elapsed_sec': 10.0, 'n_delivered': 1})
    wb = openpyxl.Workbook()
    wb.active['A1'] = 1
    wb.save(ep / 'target.xlsx')
    return ep


def make_empty_run(root: Path, run: str = 'batch-1-00002-d1'):
    (root / 'runs' / run).mkdir(parents=True)
    (root / 'runs' / run / 'broken').symlink_to('/nonexistent/path')


def make_blocked_run(root: Path, run: str = 'batch-1-00003-d2'):
    diag = root / 'runs' / run / 'spreadsheet' / 'task-00003' / 'ep01' / 'diagnostics'
    diag.mkdir(parents=True)
    write(diag / 'leak-check.blocked-123.json', {'error': 'timeout'})


def make_batch(root: Path):
    batch = root / 'batch'
    write(batch / 'status.json', {
        'task-00001': {'stage': 'design_done', 'design_run': 'batch-1-00001-d1',
                       'leak': 'pending', 'gate': 'pending'},
        'task-00003': {'design_run': 'batch-1-00003-d2'},
    })
    (batch / 'RESULTS.md').write_text('# results\n')
    (batch / '.env').write_text('SECRET=1\n')
    return batch


class IndexTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        make_full_episode(self.root)
        make_empty_run(self.root)
        make_blocked_run(self.root)
        self.batch = make_batch(self.root)

    def tearDown(self):
        self.tmp.cleanup()

    def test_index_derives_stages_and_keeps_partial_attempts(self):
        idx = build_index.build(self.root / 'runs', self.batch)
        self.assertEqual(idx['schema_version'], 1)
        self.assertEqual(len(idx['tasks']), 3)
        self.assertEqual(len(idx['runs']), 3)
        by_task = {t['id']: t for t in idx['tasks']}

        full = by_task['task-00001']['attempts'][0]
        self.assertEqual(full['stages']['route']['route'], 'rl')
        self.assertAlmostEqual(full['stages']['trainee']['pass_rate'], 0.5)
        self.assertEqual(full['stages']['leak']['state'], 'accepted')
        self.assertEqual(full['stages']['gate']['state'], 'accepted')
        self.assertEqual(full['last_stage'], 'route')
        self.assertTrue(full['current'])
        self.assertIn('diagnostics/route.json', full['files']['route'])

        empty = by_task['task-00002']['attempts'][0]
        self.assertEqual(empty['last_stage'], None)
        self.assertTrue(empty['current'])          # fallback when design_run is absent
        self.assertTrue(any('no episode' in w for w in empty['warnings']))

        blocked = by_task['task-00003']['attempts'][0]
        self.assertEqual(blocked['stages']['leak']['state'], 'blocked')

    def test_summary_counts(self):
        idx = build_index.build(self.root / 'runs', self.batch)
        s = idx['batch']['summary']
        self.assertEqual(s['n_tasks'], 3)
        self.assertEqual(s['n_runs'], 3)
        self.assertEqual(s['routes'], {'rl': 1})
        self.assertEqual(s['pass_rates'], [0.5])
        self.assertEqual(s['leak']['accepted'], 1)
        self.assertEqual(s['leak']['blocked'], 1)

    def test_broken_symlink_does_not_break_scan(self):
        idx = build_index.build(self.root / 'runs', self.batch)
        empty = [r for r in idx['runs'] if r['run_id'] == 'batch-1-00002-d1'][0]
        self.assertEqual(empty['n_files'], 0)


class ApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.tmp.name)
        ep = make_full_episode(cls.root)
        make_empty_run(cls.root)
        make_blocked_run(cls.root)
        cls.batch = make_batch(cls.root)
        cls.out = cls.root / 'index.json'
        idx = build_index.build(cls.root / 'runs', cls.batch)
        cls.out.write_text(json.dumps(idx))

        lw = openpyxl.Workbook()
        lw.active['A1'] = 1
        lw.active['B1'] = 'x'
        lw.active['A2'] = 2
        lw.save(ep / 'left.xlsx')
        rw = openpyxl.Workbook()
        rw.active['A1'] = 1
        rw.active['B1'] = 'y'
        rw.active['A2'] = 3
        rw.save(ep / 'right.xlsx')

        os.environ['VIEWER_RUNS'] = str(cls.root / 'runs')
        os.environ['VIEWER_BATCH'] = str(cls.batch)
        os.environ['VIEWER_INDEX'] = str(cls.out)
        import main as server_main
        importlib.reload(server_main)
        cls.main = server_main
        from fastapi.testclient import TestClient
        cls.client = TestClient(server_main.app)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_health_and_index(self):
        self.assertEqual(self.client.get('/healthz').status_code, 200)
        r = self.client.get('/api/index')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()['batch']['id'], 'batch')

    def test_json_pointer(self):
        rel = ('batch-1-00001-d1/spreadsheet/task-00001/ep01/'
               'diagnostics/route.json')
        r = self.client.get('/api/json', params={'root': 'runs', 'path': rel,
                                                 'pointer': '/route'})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()['value'], 'rl')
        bad = self.client.get('/api/json', params={'root': 'runs', 'path': rel,
                                                   'pointer': '/nope'})
        self.assertEqual(bad.status_code, 404)

    def test_xlsx_and_diff(self):
        base = ('batch-1-00001-d1/spreadsheet/task-00001/ep01/')
        r = self.client.get('/api/xlsx', params={
            'root': 'runs', 'path': base + 'left.xlsx'})
        self.assertEqual(r.status_code, 200)
        body = r.json()
        self.assertEqual(body['sheets'], ['Sheet'])
        self.assertEqual(body['rows'][0][:2], [1, 'x'])
        d = self.client.get('/api/xlsx/diff', params={
            'root': 'runs', 'left': base + 'left.xlsx',
            'right': base + 'right.xlsx'})
        self.assertEqual(d.status_code, 200)
        self.assertEqual(d.json()['changed_cells'], 2)

    def test_text_and_tree(self):
        r = self.client.get('/api/text', params={'root': 'batch', 'path': 'RESULTS.md'})
        self.assertEqual(r.status_code, 200)
        self.assertIn('# results', r.text)
        t = self.client.get('/api/tree', params={
            'root': 'runs', 'path': 'batch-1-00001-d1', 'depth': 1})
        self.assertEqual(t.status_code, 200)
        self.assertEqual(t.json()['entries'][0]['name'], 'spreadsheet')

    def test_traversal_and_denied_paths(self):
        self.assertEqual(self.client.get('/api/text', params={
            'root': 'batch', 'path': '../../etc/passwd'}).status_code, 403)
        self.assertEqual(self.client.get('/api/text', params={
            'root': 'batch', 'path': '.env'}).status_code, 403)
        self.assertEqual(self.client.get('/api/text', params={
            'root': 'runs', 'path': '/etc/passwd'}).status_code, 403)


if __name__ == '__main__':
    unittest.main()
