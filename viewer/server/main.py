#!/usr/bin/env python3
"""Read-only HTTP API and static host for the batch viewer.

Every endpoint is a read. Paths are resolved under one of two configured roots
and checked against an explicit deny list; there is no write surface and no
authentication because the process only ever binds to 127.0.0.1.

Environment:

    VIEWER_RUNS    runs root      (default: <repo>/runs)
    VIEWER_BATCH   batch root     (default: <repo>/batch)
    VIEWER_INDEX   index.json     (default: <repo>/viewer/data/index.json)
    VIEWER_WEB     static web dir (default: <repo>/viewer/web)
"""
from __future__ import annotations

import hashlib
import json
import os
from datetime import date, datetime
from pathlib import Path
from typing import Any

from fastapi import FastAPI, HTTPException, Query
from fastapi.responses import FileResponse, JSONResponse, PlainTextResponse
from fastapi.staticfiles import StaticFiles

REPO = Path(__file__).resolve().parents[2]
DEFAULT_BATCH = REPO / 'batch'

ROOTS = {
    'runs': Path(os.environ.get('VIEWER_RUNS') or REPO / 'runs').resolve(),
    'batch': Path(os.environ.get('VIEWER_BATCH') or DEFAULT_BATCH).resolve(),
}
INDEX_PATH = Path(os.environ.get('VIEWER_INDEX') or
                  REPO / 'viewer' / 'data' / 'index.json').resolve()
WEB_ROOT = Path(os.environ.get('VIEWER_WEB') or
                REPO / 'viewer' / 'web').resolve()

DENY_PARTS = {'.git', '.venv', 'node_modules', '__pycache__', 'sessions',
              '.claude', '.mypy_cache', '.pytest_cache'}
DENY_NAMES = {'.env', '.claude.json', 'id_rsa', 'id_ed25519', 'credentials',
              '.netrc', '.npmrc', '.pypirc'}
DENY_SUFFIXES = {'.key', '.pem', '.p12', '.pfx', '.keystore'}

MAX_JSON = 8 * 1024 * 1024
MAX_TEXT = 2 * 1024 * 1024
MAX_FILE = 20 * 1024 * 1024

app = FastAPI(title='ADB batch viewer', docs_url=None, redoc_url=None)


@app.middleware('http')
async def _headers(request, call_next):
    response = await call_next(request)
    if request.url.path == '/' or request.url.path.startswith('/static'):
        response.headers['Cache-Control'] = 'no-cache'
    response.headers['X-Content-Type-Options'] = 'nosniff'
    return response
if WEB_ROOT.is_dir():
    app.mount('/static', StaticFiles(directory=str(WEB_ROOT)), name='static')

_index_cache: dict[str, Any] = {'mtime': None, 'data': None}


def _root(key: str) -> Path:
    if key not in ROOTS:
        raise HTTPException(400, f'unknown root {key!r}')
    return ROOTS[key]


def _denied(rel: Path) -> bool:
    if any(part in DENY_PARTS for part in rel.parts):
        return True
    if rel.name in DENY_NAMES:
        return True
    if rel.name.startswith('.env'):
        return True
    return rel.suffix.lower() in DENY_SUFFIXES


def _resolve(key: str, rel: str, *, must_exist: bool = True) -> Path:
    root = _root(key)
    raw = (rel or '').strip()
    if raw.startswith('/'):
        raise HTTPException(403, 'absolute paths are not served')
    if '\x00' in raw:
        raise HTTPException(400, 'bad path')
    candidate = Path(raw)
    if candidate.is_absolute() or any(part == '..' for part in candidate.parts):
        raise HTTPException(403, 'path escapes root')
    if _denied(candidate):
        raise HTTPException(403, 'path is not served')
    path = (root / candidate).resolve()
    if path != root and root not in path.parents:
        raise HTTPException(403, 'path escapes root')
    if must_exist and not path.exists():
        raise HTTPException(404, f'not found: {rel}')
    return path


def _read_json_file(path: Path, limit: int = MAX_JSON) -> Any:
    if path.stat().st_size > limit:
        raise HTTPException(413, 'file is too large to parse')
    try:
        return json.loads(path.read_text(errors='replace'))
    except (OSError, ValueError) as exc:
        raise HTTPException(422, f'unreadable JSON: {exc}') from exc


def _pointer(obj: Any, ptr: str) -> Any:
    if not ptr:
        return obj
    if not ptr.startswith('/'):
        raise HTTPException(400, 'pointer must start with /')
    cur = obj
    for raw in ptr.split('/')[1:]:
        token = raw.replace('~1', '/').replace('~0', '~')
        try:
            if isinstance(cur, list):
                cur = cur[int(token)]
            elif isinstance(cur, dict):
                cur = cur[token]
            else:
                raise KeyError(token)
        except (KeyError, IndexError, ValueError, TypeError) as exc:
            raise HTTPException(404, f'pointer not found: {ptr}') from exc
    return cur


_PATH_TAGS = sorted(
    ((str(root) + os.sep, f'{key}:') for key, root in ROOTS.items()),
    key=lambda pair: len(pair[0]), reverse=True)


def _jsonable(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Path):
        value = str(value)
    if isinstance(value, str):
        for prefix, tag in _PATH_TAGS:
            if value.startswith(prefix):
                return tag + value[len(prefix):]
        return value
    if isinstance(value, dict):
        return {str(k): _jsonable(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_jsonable(v) for v in value]
    if isinstance(value, (int, float, bool)) or value is None:
        return value
    return str(value)


def _stamp(path: Path, rel: str, root: str) -> dict:
    st = path.stat()
    return {
        'root': root,
        'path': rel,
        'bytes': st.st_size,
        'mtime': datetime.fromtimestamp(st.st_mtime).isoformat(),
    }


@app.get('/healthz')
def healthz():
    return {
        'ok': True,
        'index': INDEX_PATH.is_file(),
        'roots': {k: str(v) for k, v in ROOTS.items()},
        'web': str(WEB_ROOT),
    }


@app.get('/api/index')
def api_index():
    if not INDEX_PATH.is_file():
        raise HTTPException(503, 'index not built')
    mtime = INDEX_PATH.stat().st_mtime
    if _index_cache['mtime'] != mtime:
        _index_cache['data'] = _read_json_file(INDEX_PATH)
        _index_cache['mtime'] = mtime
    return JSONResponse(_index_cache['data'],
                        headers={'Cache-Control': 'no-cache'})


@app.get('/api/json')
def api_json(root: str = 'runs', path: str = '', pointer: str = ''):
    p = _resolve(root, path)
    if not p.is_file():
        raise HTTPException(404, 'not a file')
    data = _read_json_file(p)
    value = _pointer(data, pointer)
    body = _stamp(p, path, root)
    body['value'] = _jsonable(value)
    body['sha256'] = hashlib.sha256(p.read_bytes()).hexdigest()
    return JSONResponse(body, headers={'Cache-Control': 'no-cache'})


@app.get('/api/tree')
def api_tree(root: str = 'runs', path: str = '', depth: int = 1):
    p = _resolve(root, path)
    if not p.is_dir():
        raise HTTPException(400, 'not a directory')
    depth = max(1, min(depth, 3))

    def walk(d: Path, level: int) -> list[dict]:
        out = []
        try:
            entries = sorted(d.iterdir(), key=lambda x: (not x.is_dir(), x.name.lower()))
        except OSError:
            return out
        for entry in entries:
            try:
                rel = entry.relative_to(p)
            except ValueError:
                rel = Path(entry.name)
            if _denied(rel):
                continue
            try:
                is_dir = entry.is_dir()
                st = entry.stat()
            except OSError:
                continue
            node = {'name': entry.name, 'type': 'dir' if is_dir else 'file',
                    'bytes': 0 if is_dir else st.st_size,
                    'mtime': datetime.fromtimestamp(st.st_mtime).isoformat(),
                    'ext': '' if is_dir else entry.suffix.lower()}
            if is_dir and level < depth:
                node['children'] = walk(entry, level + 1)
            out.append(node)
        return out

    return {'root': root, 'path': path, 'depth': depth, 'entries': walk(p, 1)}


@app.get('/api/text')
def api_text(root: str = 'runs', path: str = ''):
    p = _resolve(root, path)
    if not p.is_file():
        raise HTTPException(404, 'not a file')
    if p.stat().st_size > MAX_TEXT:
        raise HTTPException(413, 'text file is too large')
    raw = p.read_bytes()
    if b'\x00' in raw[:4096]:
        raise HTTPException(415, 'binary file')
    return PlainTextResponse(raw.decode('utf-8', errors='replace'),
                             headers={'Cache-Control': 'no-cache'})


@app.get('/api/file')
def api_file(root: str = 'runs', path: str = ''):
    p = _resolve(root, path)
    if not p.is_file():
        raise HTTPException(404, 'not a file')
    if p.stat().st_size > MAX_FILE:
        raise HTTPException(413, 'file is too large')
    return FileResponse(p, filename=p.name)


def _cell(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


@app.get('/api/xlsx')
def api_xlsx(root: str = 'runs', path: str = '', sheet: str = '',
             mode: str = 'values', max_rows: int = 200, max_cols: int = 50):
    import openpyxl

    p = _resolve(root, path)
    if not p.is_file() or p.suffix.lower() not in ('.xlsx', '.xlsm'):
        raise HTTPException(400, 'not an xlsx file')
    if p.stat().st_size > MAX_FILE:
        raise HTTPException(413, 'workbook is too large')
    max_rows = max(1, min(max_rows, 1000))
    max_cols = max(1, min(max_cols, 200))
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=(mode != 'formulas'))
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types
        raise HTTPException(422, f'cannot read workbook: {exc}') from exc
    try:
        names = list(wb.sheetnames)
        if not names:
            raise HTTPException(422, 'workbook has no sheets')
        name = sheet or names[0]
        if name not in names:
            raise HTTPException(404, f'no sheet {name!r}')
        ws = wb[name]
        rows: list[list[Any]] = []
        for r, row in enumerate(ws.iter_rows(max_row=max_rows, max_col=max_cols,
                                             values_only=True), start=1):
            rows.append([_cell(v) for v in row])
        body = {
            'root': root, 'path': path, 'mode': mode, 'sheet': name,
            'sheets': names,
            'max_row': ws.max_row, 'max_column': ws.max_column,
            'truncated': bool(ws.max_row and ws.max_row > max_rows),
        }
        body.update(_stamp(p, path, root))
        body['rows'] = rows
        return JSONResponse(_jsonable(body), headers={'Cache-Control': 'no-cache'})
    finally:
        wb.close()


@app.get('/api/xlsx/diff')
def api_xlsx_diff(root: str = 'runs', left: str = '', right: str = '',
                  max_cells: int = 500, max_rows: int = 200, max_cols: int = 60):
    import openpyxl

    lp = _resolve(root, left)
    rp = _resolve(root, right)
    for p in (lp, rp):
        if not p.is_file() or p.suffix.lower() not in ('.xlsx', '.xlsm'):
            raise HTTPException(400, 'both paths must be xlsx files')
        if p.stat().st_size > MAX_FILE:
            raise HTTPException(413, 'workbook is too large')
    max_cells = max(1, min(max_cells, 2000))
    try:
        lw = openpyxl.load_workbook(lp, read_only=True, data_only=True)
        rw = openpyxl.load_workbook(rp, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(422, f'cannot read workbook: {exc}') from exc
    try:
        sheets = sorted(set(lw.sheetnames) | set(rw.sheetnames))
        changes: list[dict] = []
        per_sheet: dict[str, int] = {}
        for name in sheets:
            ls = lw[name] if name in lw.sheetnames else None
            rs = rw[name] if name in rw.sheetnames else None
            n = 0
            if ls is None or rs is None:
                per_sheet[name] = -1   # sheet only on one side
                continue
            rows = min(ls.max_row or 0, rs.max_row or 0, max_rows)
            cols = min(ls.max_column or 0, rs.max_column or 0, max_cols)
            lgrid = list(ls.iter_rows(max_row=rows, max_col=cols, values_only=True))
            rgrid = list(rs.iter_rows(max_row=rows, max_col=cols, values_only=True))
            for ri in range(rows):
                lrow = lgrid[ri] if ri < len(lgrid) else ()
                rrow = rgrid[ri] if ri < len(rgrid) else ()
                for ci in range(cols):
                    lv = lrow[ci] if ci < len(lrow) else None
                    rv = rrow[ci] if ci < len(rrow) else None
                    if lv == rv:
                        continue
                    n += 1
                    if len(changes) < max_cells:
                        changes.append({
                            'sheet': name,
                            'cell': f'{openpyxl.utils.get_column_letter(ci + 1)}{ri + 1}',
                            'left': _cell(lv), 'right': _cell(rv),
                        })
            per_sheet[name] = n
        return {
            'root': root, 'left': left, 'right': right,
            'changed_cells': sum(v for v in per_sheet.values() if v > 0),
            'per_sheet': per_sheet,
            'truncated': len(changes) >= max_cells,
            'changes': changes,
        }
    finally:
        lw.close()
        rw.close()


@app.get('/')
def home():
    index = WEB_ROOT / 'index.html'
    if not index.is_file():
        raise HTTPException(503, 'web assets missing')
    return FileResponse(index)
